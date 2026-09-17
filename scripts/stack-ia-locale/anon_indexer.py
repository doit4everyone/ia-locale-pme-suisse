"""
indexer.py : indexation incrémentale de documents locaux vers Qdrant
avec détection automatique de l'organisation propriétaire.

Indexation incrémentale :
  Pour chaque fichier, le content_hash du texte extrait est comparé à celui
  stocké dans Qdrant, avec les paramètres d'indexation (modèle d'embedding,
  taille et chevauchement des chunks, seuil minimal, version du chunker).
  Si tout est identique, le fichier est sauté sans appel à Ollama.
  Si le fichier a changé, les ACL déjà présentes (autorises, interdits)
  sont reportées sur les nouveaux chunks : il n'y a plus de fenêtre sans
  ACL entre indexer.py et acl_resolver.py pour un fichier modifié.
  Un fichier NOUVEAU reste sans ACL, donc invisible pour les utilisateurs
  filtrés, jusqu'au passage de acl_resolver.py (comportement fail-closed).

  Les nouveaux chunks sont écrits par upsert (mêmes IDs), puis les chunks
  excédentaires de l'ancienne version sont supprimés. Le document reste
  interrogeable pendant toute la réindexation.

  Premier passage après mise à jour du script : les chunks existants ne
  portent pas encore les champs chunk_size, chunk_overlap, min_chunk_words
  et chunker_version. Tous les fichiers sont donc réindexés une fois.

Cascade de détection de l'organisation (par ordre de priorité) :
  1. --org passé en argument CLI
  2. Pattern "Client : X" dans le contenu du document
  3. Nom du dossier parent (si structure hiérarchique par client)
  4. Mots-clés de document interne → ORG_OWNER
  5. Quarantaine : document non identifiable, rapport généré

Note sur le niveau 5 :
  La classification LLM a été délibérément exclue de la cascade.
  Un LLM qui hallucine peut attribuer un document à la mauvaise
  organisation, et cette erreur serait stockée comme un fait dans
  le payload. Mieux vaut une quarantaine avec décision humaine qu'une
  attribution incorrecte avec assurance.

Note sur org_name vs autorises[] :
  org_name est une donnée DÉDUITE par heuristique. Elle est destinée
  au confort de navigation et aux facettes de recherche uniquement.
  Elle ne doit jamais servir de critère d'accès.

  Le cloisonnement repose sur autorises[] et interdits[], lus depuis les
  ACL NTFS par acl_resolver.py. Ce sont des données CONSTATÉES.

Usage :
  python indexer.py --corpus /chemin/vers/dossier
  python indexer.py --corpus /chemin/vers/dossier --org "Baumont SA"
  python indexer.py --corpus /chemin/vers/dossier --force
  python indexer.py --corpus /chemin/vers/dossier --reset
  python indexer.py --corpus /chemin/vers/dossier --rapport /tmp/rapport.json

Variables d'environnement :
  ORG_OWNER                : nom de l'organisation qui déploie la solution
  OLLAMA_URL               : URL Ollama pour les embeddings
  EMBED_MODEL              : modèle d'embedding (doit correspondre à main.py)
  QDRANT_URL               : URL Qdrant
  QDRANT_COLLECTION        : collection du corpus entreprise
  DOCUMENTATION_COLLECTION : collection de la documentation technique
  DOCUMENTATION_PATHS      : dossiers racine routés vers la documentation
  CHUNK_SIZE               : taille cible d'un chunk, en mots
  CHUNK_OVERLAP            : chevauchement entre chunks, en mots
  MIN_CHUNK_WORDS          : taille minimale d'un chunk autonome, en mots (défaut 8)

Codes de sortie :
  0 : indexation terminée sans erreur
  1 : erreur bloquante (corpus, Qdrant, Ollama, configuration)
  2 : indexation terminée, mais au moins un fichier en erreur (voir rapport)
"""

import argparse
import fnmatch
import hashlib
import json
import os
import re
import sys
import httpx
from datetime import datetime, timezone
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance, VectorParams, PointStruct,
    Filter, FieldCondition, MatchValue, Range, PayloadSchemaType
)

# Les trois bibliothèques d'extraction sont testées séparément :
# l'absence de l'une ne doit pas désactiver les deux autres.
try:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# ─────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────

OLLAMA_URL    = os.getenv("OLLAMA_URL",        "http://<IP-HOTE-OLLAMA>:11434")
EMBED_MODEL   = os.getenv("EMBED_MODEL",       "nomic-embed-text")
QDRANT_URL    = os.getenv("QDRANT_URL",        "http://localhost:6333")
COLLECTION    = os.getenv("QDRANT_COLLECTION", "documents")
DOCUMENTATION_COLLECTION = os.getenv("DOCUMENTATION_COLLECTION", "documentation")
DOCUMENTATION_PATHS = [
    p.strip() for p in
    os.getenv("DOCUMENTATION_PATHS", "DOIT4EVERYONE").split(",")
    if p.strip()
]
ORG_OWNER       = os.getenv("ORG_OWNER",       "Organisation interne")
CHUNK_SIZE      = int(os.getenv("CHUNK_SIZE",      "150"))
CHUNK_OVERLAP   = int(os.getenv("CHUNK_OVERLAP",   "20"))
# Défaut aligné sur la stack validée (.env et main.py) : une valeur différente
# déclencherait une réindexation complète, min_chunk_words étant comparé.
MIN_CHUNK_WORDS = int(os.getenv("MIN_CHUNK_WORDS", "8"))

# Version de l'algorithme de découpage. À incrémenter à chaque modification
# de chunk_blocks() : les fichiers indexés avec une autre version sont
# réindexés automatiquement au passage suivant.
CHUNKER_VERSION = 2

# Extensions indexées. acl_resolver.py doit utiliser EXACTEMENT la même liste,
# sinon ses chunks sont considérés comme orphelins et supprimés.
SUPPORTED_EXTENSIONS = ('.docx', '.pdf', '.pptx', '.txt', '.md')

# Dimensions par modèle d'embedding
MODEL_DIMS = {
    "nomic-embed-text": 768,
    "nomic-embed-text:latest": 768,
    "bge-m3": 1024,
    "bge-m3:latest": 1024,
    "multilingual-e5-base": 768,
    "multilingual-e5-large": 1024,
    "intfloat/multilingual-e5-base": 768,
}

# Patterns de détection du client dans le contenu (niveau 2)
# Contraintes appliquées après capture : voir _valider_org()
CLIENT_PATTERNS = [
    r'Client\s*:\s*([A-ZÀÂÄÉÈÊËÏÎÔÖÙÛÜÇ][^\n,]{3,50})',
    r'client\s*:\s*([A-ZÀÂÄÉÈÊËÏÎÔÖÙÛÜÇ][^\n,]{3,50})',
    r'Entre\s+.+?et\s+([A-ZÀÂÄÉÈÊËÏÎÔÖÙÛÜÇ][^\n,]{3,50}?)\s*(?:\(client\)|,|$)',
    r'pour le compte de\s+([A-ZÀÂÄÉÈÊËÏÎÔÖÙÛÜÇ][^\n,]{3,50})',
    r'destiné à\s+([A-ZÀÂÄÉÈÊËÏÎÔÖÙÛÜÇ][^\n,]{3,50})',
]

# Terminaisons légales reconnues
LEGAL_SUFFIXES = {
    'sa', 'sarl', 'sàrl', 'ag', 'gmbh', 'snc', 'sas', 'inc',
    'associés', 'associes', 'industries', 'fiduciaire', 'avocats',
    'notaires', 'conseil', 'services', 'group', 'groupe',
}

# Fichiers à exclure (bruit)
EXCLUDE_FILE_PATTERNS = [
    '~$*', '*.tmp', '*.lnk', '.DS_Store', 'Thumbs.db', 'desktop.ini',
]

# Répertoires à exclure : élagués pendant le parcours, leur contenu
# n'est jamais lu. Un pattern de fichier ne suffit pas pour un dossier.
EXCLUDE_DIR_PATTERNS = [
    'DfsrPrivate',                 # Réplication DFS, métadonnées internes
    'System Volume Information',
    '$RECYCLE.BIN',
]

# Mots-clés de documents internes (niveau 4)
# Ces mots-clés doivent être représentatifs de l'organisation qui DÉPLOIE
# la solution, pas de ses clients.
INTERNAL_KEYWORDS = [
    'politique interne',
    'procédure interne',
    'usage interne',
    'plan stratégique',
    'budget prévisionnel',
    'procès-verbal',
    'runbook',
    'politique de sécurité informatique',
    'procédure de déploiement',
    'politique rh',
    'politique des ressources humaines',
    'registre du personnel',
    'registre de',
]

# Champs ACL écrits par acl_resolver.py, reportés lors d'une réindexation
ACL_FIELDS = ("autorises", "interdits", "acl_updated_at")


def valider_configuration() -> None:
    """Arrête le script si les paramètres de découpage sont incohérents."""
    erreurs = []
    if CHUNK_SIZE < 1:
        erreurs.append(f"CHUNK_SIZE doit être positif (valeur : {CHUNK_SIZE})")
    if CHUNK_OVERLAP < 0 or CHUNK_OVERLAP >= CHUNK_SIZE:
        erreurs.append(
            f"CHUNK_OVERLAP doit être compris entre 0 et CHUNK_SIZE - 1 "
            f"(valeur : {CHUNK_OVERLAP}, CHUNK_SIZE : {CHUNK_SIZE})"
        )
    if MIN_CHUNK_WORDS < 1 or MIN_CHUNK_WORDS > CHUNK_SIZE:
        erreurs.append(
            f"MIN_CHUNK_WORDS doit être compris entre 1 et CHUNK_SIZE "
            f"(valeur : {MIN_CHUNK_WORDS})"
        )
    if erreurs:
        for e in erreurs:
            print(f"Erreur de configuration : {e}")
        sys.exit(1)

# ─────────────────────────────────────────
# Extraction de texte
# ─────────────────────────────────────────

def extract_docx_structured(path: str) -> list[str]:
    if not DOCX_AVAILABLE:
        return extract_docx_fallback(path)
    try:
        doc = DocxDocument(path)
        blocs = []
        for element in doc.element.body:
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag
            if tag == 'p':
                style = element.find(f'.//{qn("w:pStyle")}')
                style_val = ''
                if style is not None:
                    for attr in style.attrib:
                        if 'val' in attr:
                            style_val = style.attrib[attr]
                            break
                text = ''.join(
                    node.text or '' for node in element.iter()
                    if node.tag.split('}')[-1] == 't'
                ).strip()
                if text:
                    if 'Heading' in style_val or 'heading' in style_val:
                        blocs.append(f"## {text}")
                    else:
                        blocs.append(text)
            elif tag == 'tbl':
                for row in element.iter():
                    if row.tag.split('}')[-1] == 'tr':
                        cells = []
                        for cell in row.iter():
                            if cell.tag.split('}')[-1] == 'tc':
                                cell_text = ''.join(
                                    node.text or '' for node in cell.iter()
                                    if node.tag.split('}')[-1] == 't'
                                ).strip()
                                if cell_text:
                                    cells.append(cell_text)
                        if cells:
                            blocs.append(' | '.join(cells))
        return blocs
    except Exception as e:
        print(f"  Erreur python-docx : {e}, fallback XML")
        return extract_docx_fallback(path)


def extract_docx_fallback(path: str) -> list[str]:
    import zipfile
    try:
        with zipfile.ZipFile(path, 'r') as z:
            with z.open('word/document.xml') as f:
                xml = f.read().decode('utf-8', errors='ignore')
        text = re.sub(r'<[^>]+>', ' ', xml)
        text = re.sub(r'\s+', ' ', text).strip()
        return [text] if text else []
    except Exception as e:
        print(f"  Erreur lecture {path} : {e}")
        return []


def extract_pdf(path: str) -> list[str]:
    """Extrait le texte d'un PDF page par page via pdfplumber."""
    if not PDF_AVAILABLE:
        print("  pdfplumber absent : PDF ignoré")
        return []
    blocks = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text and text.strip():
                    blocks.append(text.strip())
        if not blocks:
            print(f"  PDF vide ou scanné (aucun texte extractible) : {path}")
        return blocks
    except Exception as e:
        print(f"  Erreur PDF {path} : {e}")
        return []


def extract_pptx(path: str) -> list[str]:
    """Extrait le texte d'un PowerPoint, une slide par bloc."""
    if not PPTX_AVAILABLE:
        print("  python-pptx absent : PPTX ignoré")
        return []
    blocks = []
    try:
        prs = Presentation(path)
        for slide in prs.slides:
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                blocks.append("\n".join(slide_texts))
        if not blocks:
            print(f"  PPTX vide (aucun texte extractible) : {path}")
        return blocks
    except Exception as e:
        print(f"  Erreur PPTX {path} : {e}")
        return []


def extract_txt(path: str) -> list[str]:
    try:
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        return [p.strip() for p in content.split('\n\n') if p.strip()]
    except Exception as e:
        print(f"  Erreur lecture {path} : {e}")
        return []


def extract_blocks(path: str) -> list[str]:
    ext = os.path.splitext(path)[1].lower()
    if ext == '.docx':
        return extract_docx_structured(path)
    elif ext == '.pdf':
        return extract_pdf(path)
    elif ext == '.pptx':
        return extract_pptx(path)
    elif ext in ('.txt', '.md'):
        return extract_txt(path)
    return []

# ─────────────────────────────────────────
# Chunking (version 2)
# ─────────────────────────────────────────

def chunk_blocks(
    blocks: list[str],
    chunk_size: int | None = None,
    overlap: int | None = None,
    min_words: int | None = None,
) -> list[str]:
    """
    Regroupe les blocs en chunks d'environ chunk_size mots.

    Règles :
    - Un titre (bloc "## ...") ferme le chunk en cours et ouvre une section.
      Le titre est répété en tête de chaque chunk de sa section.
    - Un bloc qui dépasse chunk_size est découpé en fenêtres glissantes.
      Le contenu en attente est placé DEVANT ce bloc, l'ordre du texte
      est conservé (correction de la version 1).
    - Un reste inférieur à min_words n'est jamais perdu : il est rattaché
      au chunk précédent, ou reporté en tête du chunk suivant s'il n'y a
      pas encore de chunk (correction de la version 1).
    - Un document entier plus court que min_words produit un chunk unique.
    """
    chunk_size = chunk_size if chunk_size is not None else CHUNK_SIZE
    overlap = overlap if overlap is not None else CHUNK_OVERLAP
    min_words = min_words if min_words is not None else MIN_CHUNK_WORDS

    chunks: list[str] = []
    titres_chunks: list[str] = []   # titre de section de chaque chunk émis
    section: list[str] = []         # mots du titre de la section courante
    contenu: list[str] = []         # mots du chunk en cours, hors titre
    report: list[str] = []          # reste court en attente d'un premier chunk

    def emettre(mots: list[str]) -> bool:
        """Émet un chunk. Retourne True si un chunk autonome a été créé."""
        nonlocal report
        if not mots:
            return False
        titre = ' '.join(section)
        if len(mots) >= min_words:
            chunks.append(' '.join(report + section + mots))
            titres_chunks.append(titre)
            report = []
            return True
        if chunks:
            # Même section : pas besoin de répéter le titre
            ajout = mots if titres_chunks[-1] == titre else section + mots
            chunks[-1] += ' ' + ' '.join(ajout)
        else:
            report = report + section + mots
        return False

    for block in blocks:
        if block.startswith("## "):
            emettre(contenu)
            section = block.split()
            contenu = []
            continue

        mots = block.split()

        if len(contenu) + len(mots) <= chunk_size:
            contenu.extend(mots)
            continue

        if len(mots) > chunk_size:
            mots = contenu + mots
            contenu = []
            pas = chunk_size - overlap
            for j in range(0, len(mots), pas):
                emettre(mots[j:j + chunk_size])
                if j + chunk_size >= len(mots):
                    break
            continue

        # Le bloc ferait déborder le chunk en cours
        precedent = contenu
        autonome = emettre(precedent)
        chevauchement = precedent[-overlap:] if (overlap and autonome) else []
        contenu = chevauchement + mots

    emettre(contenu)

    if report:
        chunks.append(' '.join(report))

    return chunks

# ─────────────────────────────────────────
# Détection de l'organisation (cascade 5 niveaux, LLM exclu)
# ─────────────────────────────────────────

def detect_org_level1(org_arg: str) -> str | None:
    """Niveau 1 : --org passé en argument CLI."""
    return org_arg.strip() if org_arg and org_arg.strip() else None


def _valider_org(org: str) -> str | None:
    """
    Valide et nettoie une capture de nom d'organisation.
    Retourne le nom nettoyé ou None si invalide.
    """
    while org and org[-1] in '.,;:()':
        org = org[:-1].strip()
    org = re.sub(r'\s*\([^)]*\)\s*$', '', org).strip()
    if len(org) < 6:
        return None
    if re.search(r'\d', org):
        return None
    TERMES_EXCLUS = {
        'durée', 'employeur', 'employée', 'employé',
        'salarié', 'salariée', 'contrat', 'période',
    }
    if org.lower().strip() in TERMES_EXCLUS:
        return None
    last_word = org.split()[-1].lower().rstrip('.,;:')
    has_legal_suffix = last_word in LEGAL_SUFFIXES
    if not has_legal_suffix and len(org) < 25:
        return None
    if len(org) > 40 and not has_legal_suffix:
        return None
    return org


def detect_org_level2(blocks: list[str]) -> str | None:
    """Niveau 2 : pattern "Client : X" dans les 8 premiers blocs."""
    search_text = ' '.join(blocks[:8])
    for pattern in CLIENT_PATTERNS:
        match = re.search(pattern, search_text, re.IGNORECASE)
        if match:
            org = _valider_org(match.group(1).strip())
            if org:
                return org
    return None


def detect_org_level3(filepath: str, corpus_root: str) -> str | None:
    """
    Niveau 3 : nom du dossier parent, sauf si le parent est la racine.
    Ne convient que si la structure est organisée par client.
    """
    parent = os.path.dirname(filepath)
    if os.path.abspath(parent) == os.path.abspath(corpus_root):
        return None
    folder_name = os.path.basename(parent)
    org = folder_name.replace('_', ' ').replace('-', ' ').strip()
    if len(org) >= 3:
        return org
    return None


def detect_org_level4(blocks: list[str]) -> str | None:
    """Niveau 4 : mots-clés de document interne → ORG_OWNER."""
    search_text = ' '.join(blocks[:10]).lower()
    for keyword in INTERNAL_KEYWORDS:
        if keyword.lower() in search_text:
            return ORG_OWNER
    return None


def detect_organisation(
    filepath: str,
    corpus_root: str,
    blocks: list[str],
    org_arg: str = ""
) -> tuple[str, str, int]:
    """Cascade complète. Retourne (org_name, methode, niveau)."""
    org = detect_org_level1(org_arg)
    if org:
        return org, "argument CLI (--org)", 1
    org = detect_org_level2(blocks)
    if org:
        return org, "pattern 'Client :' dans le contenu", 2
    org = detect_org_level3(filepath, corpus_root)
    if org:
        return org, "nom du dossier parent", 3
    org = detect_org_level4(blocks)
    if org:
        return org, "mots-clés document interne", 4
    return "QUARANTAINE", "non identifiable", 5

# ─────────────────────────────────────────
# Embedding
# ─────────────────────────────────────────

# Client HTTP unique, réutilisé pour tous les chunks : évite d'ouvrir
# une connexion TCP par embedding.
_http_client: httpx.Client | None = None


def http_client() -> httpx.Client:
    global _http_client
    if _http_client is None:
        _http_client = httpx.Client(timeout=60)
    return _http_client


def get_embedding(text: str) -> list[float]:
    import re as _re
    # Nettoyer les suites de points répétitifs (tables des matières PDF)
    # qui font planter nomic-embed-text avec une erreur 500.
    text = _re.sub(r'[. ]{10,}', ' ', text).strip()
    r = http_client().post(
        f"{OLLAMA_URL}/api/embeddings",
        json={"model": EMBED_MODEL, "prompt": text}
    )
    r.raise_for_status()
    return r.json()["embedding"]

# ─────────────────────────────────────────
# Qdrant
# ─────────────────────────────────────────

def get_embed_dim() -> int:
    """Retourne les dimensions pour EMBED_MODEL, avec vérification."""
    dim = MODEL_DIMS.get(EMBED_MODEL)
    if dim is None:
        print(f"  Modèle '{EMBED_MODEL}' absent de MODEL_DIMS, détection automatique...")
        try:
            dim = len(get_embedding("test"))
            print(f"  Dimensions détectées : {dim}")
        except Exception as e:
            print(f"  Erreur détection dimensions : {e}")
            print(f"  Ajouter '{EMBED_MODEL}' dans MODEL_DIMS dans indexer.py")
            sys.exit(1)
    return dim


def get_collection_for_path(source_name: str) -> str:
    """
    DOCUMENTATION_PATHS (racine du partage) → DOCUMENTATION_COLLECTION.
    Tout le reste → COLLECTION (corpus entreprise).
    """
    for prefix in DOCUMENTATION_PATHS:
        if source_name.startswith(prefix + "/") or source_name.startswith(prefix + "\\"):
            return DOCUMENTATION_COLLECTION
    return COLLECTION


def init_collection(qdrant: QdrantClient, reset: bool = False) -> int:
    """
    Crée ou vérifie les deux collections Qdrant.
    Arrête si une collection existante a des dimensions incompatibles.
    """
    embed_dim = get_embed_dim()
    existing = [c.name for c in qdrant.get_collections().collections]

    for col in [COLLECTION, DOCUMENTATION_COLLECTION]:
        if reset and col in existing:
            qdrant.delete_collection(col)
            print(f"Collection '{col}' supprimée.")
            existing = [c.name for c in qdrant.get_collections().collections]

        if col not in existing:
            qdrant.create_collection(
                collection_name=col,
                vectors_config=VectorParams(size=embed_dim, distance=Distance.COSINE)
            )
            print(f"Collection '{col}' créée (dim={embed_dim}, cosine, modèle={EMBED_MODEL}).")
        else:
            info = qdrant.get_collection(col)
            existing_dim = info.config.params.vectors.size
            if existing_dim != embed_dim:
                print(f"ERREUR : collection '{col}' existante en {existing_dim} dimensions,")
                print(f"mais '{EMBED_MODEL}' produit {embed_dim} dimensions.")
                print("Utilisez --reset pour recréer la collection, ou changez EMBED_MODEL.")
                sys.exit(1)
            print(f"Collection '{col}' existante conservée (dim={existing_dim}).")

        creer_index_payload(qdrant, col)

    return embed_dim


# Index de payload. Sans index, chaque filtre sur "source" parcourt toute la
# collection : indexer.py, acl_resolver.py et main.py filtrent tous sur ce
# champ, une fois par fichier à chaque passage.
PAYLOAD_INDEXES = {
    "source": PayloadSchemaType.KEYWORD,
    "chunk_index": PayloadSchemaType.INTEGER,
}


def creer_index_payload(qdrant: QdrantClient, col: str) -> None:
    """Crée les index de payload absents. Sans effet s'ils existent déjà."""
    try:
        existants = qdrant.get_collection(col).payload_schema or {}
    except Exception:
        existants = {}
    for champ, schema in PAYLOAD_INDEXES.items():
        if champ in existants:
            continue
        qdrant.create_payload_index(
            collection_name=col,
            field_name=champ,
            field_schema=schema,
            wait=True,
        )
        print(f"Index de payload créé : '{col}'.{champ} ({schema.value})")


def filtre_source(source_name: str) -> Filter:
    return Filter(must=[
        FieldCondition(key="source", match=MatchValue(value=source_name))
    ])


def lire_etat_existant(qdrant: QdrantClient, source_name: str, collection_name: str) -> dict | None:
    """
    Retourne le payload d'un chunk existant du fichier (sans le texte),
    ou None si le fichier n'est pas encore indexé dans cette collection.
    """
    points, _ = qdrant.scroll(
        collection_name=collection_name,
        scroll_filter=filtre_source(source_name),
        limit=1,
        with_payload=[
            "content_hash", "embed_model", "chunk_size", "chunk_overlap",
            "min_chunk_words", "chunker_version", "org_name",
            *ACL_FIELDS,
        ],
        with_vectors=False,
    )
    return points[0].payload if points else None


def est_inchange(existant: dict | None, content_hash: str, org_arg: str) -> bool:
    """
    Un fichier est inchangé si son contenu ET les paramètres d'indexation
    sont identiques à ceux stockés. Avec --org, l'organisation stockée
    doit aussi correspondre, sinon on réindexe pour la réétiqueter.
    """
    if not existant:
        return False
    if org_arg and existant.get("org_name") != org_arg.strip():
        return False
    return (
        existant.get("content_hash") == content_hash
        and existant.get("embed_model") == EMBED_MODEL
        and existant.get("chunk_size") == CHUNK_SIZE
        and existant.get("chunk_overlap") == CHUNK_OVERLAP
        and existant.get("min_chunk_words") == MIN_CHUNK_WORDS
        and existant.get("chunker_version") == CHUNKER_VERSION
    )


def supprimer_chunks(qdrant: QdrantClient, source_name: str, collections: list[str]) -> None:
    """
    Supprime tous les chunks d'un fichier dans les collections indiquées.
    La suppression se fait uniquement par chemin (source). L'ancienne
    suppression par content_hash a été retirée : elle effaçait les copies
    identiques du même fichier situées dans d'autres dossiers.
    """
    for col in collections:
        qdrant.delete(collection_name=col, points_selector=filtre_source(source_name))


def supprimer_chunks_excedentaires(
    qdrant: QdrantClient, source_name: str, collection_name: str, nb_chunks: int
) -> None:
    """
    Après upsert, supprime les chunks de l'ancienne version dont l'index
    dépasse le nombre de chunks de la nouvelle version.
    """
    qdrant.delete(
        collection_name=collection_name,
        points_selector=Filter(must=[
            FieldCondition(key="source", match=MatchValue(value=source_name)),
            FieldCondition(key="chunk_index", range=Range(gte=nb_chunks)),
        ])
    )


def autre_collection(collection_name: str) -> str:
    return DOCUMENTATION_COLLECTION if collection_name == COLLECTION else COLLECTION

# ─────────────────────────────────────────
# Pipeline principal
# ─────────────────────────────────────────

def index_file(
    qdrant: QdrantClient,
    filepath: str,
    corpus_root: str,
    source_name: str,
    org_arg: str = "",
    force: bool = False,
) -> tuple[int, dict]:
    """
    Indexe un fichier si nécessaire. Retourne (nb_chunks, entrée_rapport).
    Les exceptions (Ollama, Qdrant) remontent à l'appelant, qui les
    enregistre sans interrompre le traitement des autres fichiers.
    """
    target_collection = get_collection_for_path(source_name)

    blocks = extract_blocks(filepath)
    if not blocks:
        # Fichier vidé ou devenu illisible : l'ancien contenu ne doit pas
        # rester interrogeable.
        supprimer_chunks(qdrant, source_name, [COLLECTION, DOCUMENTATION_COLLECTION])
        return 0, {"fichier": source_name, "statut": "vide"}

    content_hash = hashlib.sha256(''.join(blocks).encode()).hexdigest()[:32]
    existant = lire_etat_existant(qdrant, source_name, target_collection)

    if not force and est_inchange(existant, content_hash, org_arg):
        print(f"  Inchangé [hash: {content_hash[:8]}...], ignoré")
        return 0, {"fichier": source_name, "statut": "inchangé"}

    total_chars = sum(len(b) for b in blocks)
    motif = "nouveau" if not existant else ("forcé" if force else "modifié")
    print(f"  Blocs extraits : {len(blocks)} ({total_chars} caractères) "
          f"[hash: {content_hash[:8]}...] ({motif})")

    org_name, methode, niveau = detect_organisation(filepath, corpus_root, blocks, org_arg)
    print(f"  Organisation : '{org_name}' (niveau {niveau} : {methode})")

    if niveau == 5:
        # Le contenu a changé et n'est plus identifiable : l'ancienne version
        # est retirée plutôt que laissée interrogeable sans validation.
        supprimer_chunks(qdrant, source_name, [COLLECTION, DOCUMENTATION_COLLECTION])
        return 0, {
            "fichier": source_name,
            "statut": "quarantaine",
            "organisation": None,
            "niveau": niveau,
            "methode": methode,
            "action_requise": (
                "Organisation non identifiable automatiquement. "
                "Réindexer avec : python indexer.py --corpus <dossier> --org \"Nom Organisation\""
            )
        }

    chunks = chunk_blocks(blocks)
    print(f"  Chunks générés : {len(chunks)}")
    if not chunks:
        supprimer_chunks(qdrant, source_name, [COLLECTION, DOCUMENTATION_COLLECTION])
        return 0, {"fichier": source_name, "statut": "vide"}

    # ACL reportées depuis la version précédente du fichier
    acl_reportees = {k: existant[k] for k in ACL_FIELDS if existant and k in existant}

    # Tous les embeddings sont calculés AVANT toute écriture dans Qdrant :
    # si Ollama échoue, l'ancienne version du fichier reste intacte.
    indexed_at = datetime.now(timezone.utc).isoformat()
    source_id = hashlib.md5(source_name.encode()).hexdigest()
    points = []
    for i, chunk in enumerate(chunks):
        chunk_id = int(hashlib.md5(f"{source_name}_{i}".encode()).hexdigest()[:16], 16) % (2**63)
        print(f"    Embedding chunk {i+1}/{len(chunks)}...", end='\r')
        points.append(PointStruct(
            id=chunk_id,
            vector=get_embedding(chunk),
            payload={
                "text": chunk,
                "source": source_name,
                "source_id": source_id,
                "content_hash": content_hash,
                "org_name": org_name,
                "detection_niveau": niveau,
                "detection_methode": methode,
                "embed_model": EMBED_MODEL,
                "chunk_size": CHUNK_SIZE,
                "chunk_overlap": CHUNK_OVERLAP,
                "min_chunk_words": MIN_CHUNK_WORDS,
                "chunker_version": CHUNKER_VERSION,
                "chunk_index": i,
                "filepath": filepath,
                "indexed_at": indexed_at,
                **acl_reportees,
            }
        ))

    # Écriture sans fenêtre d'indisponibilité : les IDs sont déterministes
    # (chemin + index), l'upsert remplace les chunks existants, puis on retire
    # ceux de l'ancienne version qui n'existent plus.
    qdrant.upsert(collection_name=target_collection, points=points)
    supprimer_chunks_excedentaires(qdrant, source_name, target_collection, len(points))
    # Si DOCUMENTATION_PATHS a changé, le fichier peut avoir changé de collection
    supprimer_chunks(qdrant, source_name, [autre_collection(target_collection)])

    acl_info = "ACL reportées" if acl_reportees.get("autorises") else "sans ACL, en attente de acl_resolver.py"
    print(f"  → {len(points)} chunks indexés (org: {org_name}, {acl_info})          ")

    return len(points), {
        "fichier": source_name,
        "statut": "indexé",
        "motif": motif,
        "organisation": org_name,
        "niveau": niveau,
        "methode": methode,
        "chunks": len(points),
        "acl_reportees": bool(acl_reportees.get("autorises")),
    }


def lister_fichiers(corpus_dir: str) -> tuple[list[str], list[str], list[str]]:
    """
    Parcours récursif avec élagage des répertoires exclus.
    Retourne (fichiers_supportés, fichiers_bruit, erreurs_de_parcours).
    """
    all_files: list[str] = []
    skipped_noise: list[str] = []
    walk_errors: list[str] = []

    def on_error(err: OSError) -> None:
        walk_errors.append(f"{err.filename} : {err.strerror}")

    for root, dirs, files in os.walk(corpus_dir, onerror=on_error):
        # Élagage en place : os.walk ne descend pas dans les dossiers retirés
        dirs[:] = sorted(
            d for d in dirs
            if not any(fnmatch.fnmatch(d, pat) for pat in EXCLUDE_DIR_PATTERNS)
        )
        for f in sorted(files):
            if any(fnmatch.fnmatch(f, pat) for pat in EXCLUDE_FILE_PATTERNS):
                skipped_noise.append(f)
                continue
            if os.path.splitext(f)[1].lower() in SUPPORTED_EXTENSIONS:
                all_files.append(os.path.join(root, f))

    return all_files, skipped_noise, walk_errors


def compter_chunks(qdrant: QdrantClient, col: str) -> int:
    try:
        return qdrant.count(col).count
    except Exception:
        return 0


def index_corpus(
    corpus_dir: str,
    reset: bool = False,
    org_arg: str = "",
    rapport_path: str = "",
    force: bool = False,
) -> int:
    """Indexe le corpus. Retourne le code de sortie (0 ou 2)."""
    if not os.path.isdir(corpus_dir):
        print(f"Erreur : dossier introuvable : {corpus_dir}")
        sys.exit(1)

    valider_configuration()

    print(f"Connexion Qdrant : {QDRANT_URL}")
    qdrant_client = QdrantClient(url=QDRANT_URL)
    init_collection(qdrant_client, reset=reset)

    print(f"Vérification Ollama : {OLLAMA_URL}")
    try:
        r = http_client().get(f"{OLLAMA_URL}/api/tags", timeout=10)
        models = [m['name'] for m in r.json().get('models', [])]
        if not any(EMBED_MODEL.split(':')[0] in m for m in models):
            print(f"Erreur : modèle d'embedding '{EMBED_MODEL}' non trouvé.")
            sys.exit(1)
        print(f"Modèle d'embedding : {EMBED_MODEL} ✓")
    except Exception as e:
        print(f"Erreur connexion Ollama : {e}")
        sys.exit(1)

    if org_arg:
        print(f"Organisation forcée (niveau 1) : '{org_arg}'")
    print(f"Organisation interne par défaut (niveau 4) : '{ORG_OWNER}'")
    print("Note : org_name sert à la navigation uniquement, jamais au contrôle d'accès")
    print(f"Collection documents      : '{COLLECTION}'")
    print(f"Collection documentation  : '{DOCUMENTATION_COLLECTION}'")
    print(f"Préfixes documentation    : {DOCUMENTATION_PATHS}")
    print(f"Découpage                 : {CHUNK_SIZE} mots, chevauchement {CHUNK_OVERLAP}, "
          f"minimum {MIN_CHUNK_WORDS}, chunker v{CHUNKER_VERSION}")
    print(f"Mode                      : {'réindexation forcée' if force else 'incrémental'}")

    all_files, skipped_noise, walk_errors = lister_fichiers(corpus_dir)
    if skipped_noise:
        print(f"Fichiers bruit exclus : {len(skipped_noise)} "
              f"({', '.join(skipped_noise[:5])}{'...' if len(skipped_noise) > 5 else ''})")
    if walk_errors:
        print(f"Dossiers inaccessibles pendant le parcours : {len(walk_errors)}")
        for err in walk_errors:
            print(f"  {err}")

    if not all_files:
        print(f"Aucun fichier supporté trouvé dans {corpus_dir}")
        sys.exit(1)

    print(f"\n{len(all_files)} fichiers à examiner :\n")

    rapport = []
    total_chunks = 0

    for i, filepath in enumerate(all_files, 1):
        source_name = os.path.relpath(filepath, corpus_dir)
        print(f"\n  [{i}/{len(all_files)}] {source_name}")
        try:
            nb, entry = index_file(qdrant_client, filepath, corpus_dir, source_name, org_arg, force)
        except Exception as e:
            # L'ancienne version du fichier reste en place dans Qdrant
            print(f"\n  ERREUR : {e}")
            nb, entry = 0, {"fichier": source_name, "statut": "erreur", "erreur": str(e)}
        total_chunks += nb
        rapport.append(entry)

    indexés     = [r for r in rapport if r["statut"] == "indexé"]
    inchangés   = [r for r in rapport if r["statut"] == "inchangé"]
    quarantaine = [r for r in rapport if r["statut"] == "quarantaine"]
    vides       = [r for r in rapport if r["statut"] == "vide"]
    erreurs     = [r for r in rapport if r["statut"] == "erreur"]
    sans_acl    = [r for r in indexés if not r.get("acl_reportees")]
    count_docs  = compter_chunks(qdrant_client, COLLECTION)
    count_tech  = compter_chunks(qdrant_client, DOCUMENTATION_COLLECTION)

    print(f"\n{'='*60}")
    print("RAPPORT D'INDEXATION")
    print(f"{'='*60}")
    print(f"Fichiers examinés     : {len(all_files)}")
    print(f"Inchangés (ignorés)   : {len(inchangés)}")
    print(f"Indexés               : {len(indexés)}")
    print(f"  dont sans ACL       : {len(sans_acl)} (visibles après acl_resolver.py)")
    print(f"Fichiers vides        : {len(vides)}")
    print(f"En quarantaine        : {len(quarantaine)}")
    print(f"En erreur             : {len(erreurs)}")
    print(f"Chunks écrits         : {total_chunks}")
    print(f"Total chunks Qdrant   : {count_docs} ('{COLLECTION}') + {count_tech} ('{DOCUMENTATION_COLLECTION}')")

    if indexés:
        print("\nOrganisations détectées (fichiers indexés à ce passage) :")
        orgs: dict[str, int] = {}
        for r in indexés:
            orgs[r["organisation"]] = orgs.get(r["organisation"], 0) + 1
        for org, nb in sorted(orgs.items()):
            print(f"  {nb:3d} fichier(s) → {org}")

        print("\nMéthodes de détection :")
        methodes: dict[str, int] = {}
        for r in indexés:
            m = f"Niveau {r['niveau']} : {r['methode']}"
            methodes[m] = methodes.get(m, 0) + 1
        for m, nb in sorted(methodes.items()):
            print(f"  {nb:3d} fichier(s) → {m}")

    if quarantaine:
        print(f"\n{'!'*60}")
        print(f"DOCUMENTS EN QUARANTAINE ({len(quarantaine)}) :")
        print("Action requise : réindexer avec --org")
        print(f"{'!'*60}")
        for r in quarantaine:
            print(f"  {r['fichier']}")

    if erreurs:
        print(f"\nFICHIERS EN ERREUR ({len(erreurs)}) : ancienne version conservée dans Qdrant")
        for r in erreurs:
            print(f"  {r['fichier']} : {r['erreur']}")

    if not rapport_path:
        rapport_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            f"rapport_indexation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        )
    os.makedirs(os.path.dirname(os.path.abspath(rapport_path)), exist_ok=True)

    with open(rapport_path, 'w', encoding='utf-8') as f:
        json.dump({
            "date": datetime.now(timezone.utc).isoformat(),
            "corpus": corpus_dir,
            "collections": [COLLECTION, DOCUMENTATION_COLLECTION],
            "embed_model": EMBED_MODEL,
            "org_owner": ORG_OWNER,
            "mode": "force" if force else "incremental",
            "parametres": {
                "chunk_size": CHUNK_SIZE,
                "chunk_overlap": CHUNK_OVERLAP,
                "min_chunk_words": MIN_CHUNK_WORDS,
                "chunker_version": CHUNKER_VERSION,
            },
            "statistiques": {
                "total_fichiers": len(all_files),
                "inchangés": len(inchangés),
                "indexés": len(indexés),
                "indexés_sans_acl": len(sans_acl),
                "quarantaine": len(quarantaine),
                "vides": len(vides),
                "erreurs": len(erreurs),
                "chunks_écrits": total_chunks,
                "chunks_documents": count_docs,
                "chunks_documentation": count_tech,
            },
            "erreurs_parcours": walk_errors,
            "documents": rapport
        }, f, ensure_ascii=False, indent=2)
    print(f"\nRapport sauvegardé : {rapport_path}")

    return 2 if erreurs else 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Indexeur RAG incrémental avec détection automatique d'organisation",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Exemples :
  # Passage incrémental (défaut) : seuls les fichiers nouveaux ou modifiés
  ORG_OWNER="<Nom organisation>" python indexer.py --corpus /serveur/documents

  # Dossier d'un client spécifique
  python indexer.py --corpus /serveur/Baumont_Industries --org "Baumont Industries SA"

  # Réindexer tous les fichiers en conservant les ACL
  python indexer.py --corpus /serveur/documents --force

  # Supprimer et recréer les collections (ACL perdues, relancer acl_resolver.py)
  python indexer.py --corpus /serveur/documents --reset

  # Rapport dans un dossier dédié
  python indexer.py --corpus /serveur/documents --rapport /var/log/rag/rapport.json
        """
    )
    parser.add_argument("--corpus",  required=True, help="Dossier de documents (récursif)")
    parser.add_argument("--org",     default="",    help="Forcer l'organisation (niveau 1)")
    parser.add_argument("--force",   action="store_true",
                        help="Réindexer tous les fichiers, même inchangés (ACL conservées)")
    parser.add_argument("--reset",   action="store_true",
                        help="Supprimer et recréer les collections (ACL perdues)")
    parser.add_argument("--rapport", default="",    help="Chemin du rapport JSON (défaut : dossier du script)")
    args = parser.parse_args()

    code = index_corpus(
        args.corpus,
        reset=args.reset,
        org_arg=args.org,
        rapport_path=args.rapport,
        force=args.force,
    )
    sys.exit(code)
